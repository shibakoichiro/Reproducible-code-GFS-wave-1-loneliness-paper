"""
Rebuild Supplementary Figures S1-S64 (country-specific forest plots) with
- no title inside the image (the title is a text paragraph in Word),
- a legend under each figure (as required for Supplementary Figures),
- x-axis labels that describe estimates as means / regression coefficients.

Input : either the analysis output folder (output/, containing tbl3.output.csv
        from 01_sociodemographic_variation.R and tbl2.output.csv from
        02_childhood_correlates.R) or the Appendix (.docx), in which case the
        estimates are read from Supplementary Tables S1b-S22b and S1c-S22c.
Output: <out_dir>/SI_Figures_S1-S64.docx  (one figure per page: title,
        image, legend) and <out_dir>/png/FigureS#.png (300 dpi).

Figure order follows the original Supplementary Information:
  S1-S34  : sociodemographic categories (religious affiliation and race and
            ethnicity are not shown because categories differ by country)
  S35-S61 : childhood factors
  S62-S63 : religious affiliation at age 12 (Christianity; some other
            religion), countries whose reference group is no religion
  S64     : race and ethnicity (minority groups vs plurality group)

Usage: python3 06_SI_figures.py output output     (from the reproducible_code folder)
"""
import os
import re
import sys

import docx
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC, OUT = (sys.argv[1], sys.argv[2]) if len(sys.argv) > 2 else ("output", "output")
os.makedirs(os.path.join(OUT, "png"), exist_ok=True)

# Helvetica-like font for the images, Times New Roman for Word text
for f in font_manager.findSystemFonts():
    if "texgyreheros-regular" in f.lower() or "texgyreheros-bold." in f.lower():
        font_manager.fontManager.addfont(f)
plt.rcParams["font.family"] = ["TeX Gyre Heros", "Liberation Sans", "DejaVu Sans"]
plt.rcParams["font.size"] = 9.5
FIG_W_MM = 170  # printed width in Word (A4 with 20 mm margins)
ROW_MM = 8.3     # height per country row

N_TOTAL = "202,898"

# --------------------------------------------------------------------------
# 1. Read country-specific estimates from Supplementary Tables S#b and S#c
# --------------------------------------------------------------------------
def parse_ci(s):
    m = re.match(r"\(\s*(-?[\d.]+),\s*(-?[\d.]+)\)", s.replace("−", "-"))
    return (float(m.group(1)), float(m.group(2))) if m else (None, None)


def load_from_appendix(path):
    """Read estimates from Supplementary Tables S#b/S#c of the Appendix (.docx)."""
    src = docx.Document(path)
    rows = []
    last = ""
    for el in src.element.body.iterchildren():
        if el.tag.endswith("}p"):
            t = Paragraph(el, src).text.strip()
            if t:
                last = t
            continue
        if not el.tag.endswith("}tbl"):
            continue
        m = re.match(r"Table S(\d+)([bc])\. .* for (.+)$", last)
        if not m:
            continue
        kind, country = m.group(2), m.group(3).strip()
        var = ""
        for r in Table(el, src).rows[1:]:
            cells = []
            prev = None
            for c in r.cells:
                cells.append("" if c._tc is prev else c.text.strip())
                prev = c._tc
            if cells[0].startswith("Note"):
                continue
            if cells[0]:
                var = cells[0]
            cat, est, ci = cells[1], cells[2], cells[4]
            if not est:
                continue
            lo, hi = parse_ci(ci) if ci else (None, None)
            rows.append(dict(kind=kind, country=country, variable=var, category=cat,
                             est=float(est.replace("\u2212", "-")), lo=lo, hi=hi))
    return rows


def load_from_csv(out_dir):
    """Read estimates from the analysis output (tbl3.output.csv from 01 and
    tbl2.output.csv from 02), rounded to two decimals as in the Supplementary Tables."""
    import pandas as pd
    rows = []
    d3 = pd.read_csv(os.path.join(out_dir, "tbl3.output.csv"))
    for _, r in d3.iterrows():
        if pd.isna(r["mean.pooled"]):
            continue
        ok = not pd.isna(r["vm.adj"]) and r["vm.adj"] > 1 and not pd.isna(r["ci.low"])
        rows.append(dict(kind="b", country=str(r["COUNTRY"]).strip(), variable=r["Variable"],
                         category=str(r["Category"]), est=round(float(r["mean.pooled"]), 2) + 0.0,
                         lo=round(float(r["ci.low"]), 2) if ok else None,
                         hi=round(float(r["ci.up"]), 2) if ok else None))
    d2 = pd.read_csv(os.path.join(out_dir, "tbl2.output.csv"))
    for _, r in d2.iterrows():
        lo, hi = parse_ci(str(r["CI"])) if isinstance(r["CI"], str) else (None, None)
        rows.append(dict(kind="c", country=str(r["Country"]).strip(), variable=r["Variable"],
                         category=str(r["Category"]), est=float(r["Estimate"]) + 0.0, lo=lo, hi=hi))  # + 0.0 turns -0.0 into 0.0
    # same row order as the Supplementary Tables: country, then table order
    order = {c: i for i, c in enumerate(sorted({x["country"] for x in rows}))}
    rows.sort(key=lambda x: (x["kind"], order[x["country"]]))
    return rows


rows = load_from_appendix(SRC) if SRC.lower().endswith(".docx") else load_from_csv(SRC)

COUNTRIES = sorted({r["country"] for r in rows})
assert len(COUNTRIES) == 22, COUNTRIES


def ordered_keys(kind, keep):
    seen = []
    for r in rows:
        k = (r["variable"], r["category"])
        if r["kind"] == kind and keep(r) and k not in seen:
            seen.append(k)
    return seen


demo_keys = ordered_keys(
    "b", lambda r: r["variable"] not in ("Religious affiliation", "Race and ethnicity"))
child_keys = ordered_keys(
    "c", lambda r: not r["variable"].startswith(("Religious affiliation", "Race and ethnicity")))
rel_keys = [("Religious affiliation at age 12 (Ref: No religion/atheist/agnostic)", c)
            for c in ("Christianity", "Some other religion")]
race_key = [k for k in ordered_keys("c", lambda r: r["variable"].startswith("Race and ethnicity"))]


def data_for(kind, key, var_match=None):
    out = []
    for r in rows:
        if r["kind"] != kind or r["category"] != key[1]:
            continue
        if var_match is not None:
            if not var_match(r["variable"]):
                continue
        elif r["variable"] != key[0]:
            continue
        out.append(r)
    return out


figs = []  # (kind, key, data)
for k in demo_keys:
    figs.append(("b", k, data_for("b", k)))
for k in child_keys:
    figs.append(("c", k, data_for("c", k)))
for k in rel_keys:
    figs.append(("c", k, data_for("c", k, lambda v: v.lower().startswith(
        "religious affiliation at age 12 (ref: no religion"))))
for k in race_key:
    figs.append(("c", k, data_for("c", k, lambda v: v.startswith("Race and ethnicity"))))

print(f"{len(figs)} figures ({len(demo_keys)} sociodemographic, "
      f"{len(child_keys)} childhood, {len(rel_keys)} religion, {len(race_key)} race)")
assert len(figs) == 64, len(figs)

# common x-axis range within each family
c_all = [v for kind, _, d in figs if kind == "c" for r in d for v in (r["lo"], r["hi"]) if v is not None]
C_LIM = (min(c_all) - 0.1, max(c_all) + 0.1)
B_LIM = (0, 10)

# --------------------------------------------------------------------------
# 2. Draw each forest plot (no title in the image)
# --------------------------------------------------------------------------
def draw(kind, data, path):
    data = sorted(data, key=lambda r: r["est"])
    n = len(data)
    fig_h = (18 + ROW_MM * n) / 25.4  # inches
    fig, (ax, axt) = plt.subplots(1, 2, figsize=(FIG_W_MM / 25.4, fig_h),
                                  gridspec_kw=dict(width_ratios=[2.1, 1], wspace=0.02))
    ys = range(n)
    for y in ys:
        if y % 2 == 1:
            for a in (ax, axt):
                a.axhspan(y - 0.5, y + 0.5, color="#ebebeb", zorder=0, lw=0)
    if kind == "c":
        ax.axvline(0, ls="--", color="0.3", lw=0.8, zorder=1)
    for y, r in zip(ys, data):
        if r["lo"] is not None:
            ax.plot([r["lo"], r["hi"]], [y, y], color="black", lw=1.0, zorder=2)
            lab = f"{r['est']:.2f} ({r['lo']:.2f}, {r['hi']:.2f})"
        else:
            lab = f"{r['est']:.2f} (CI not estimable)"
        ax.plot(r["est"], y, marker="s", color="black", ms=5, zorder=3)
        axt.text(0.5, y, lab, ha="center", va="center", fontsize=9.5)
    ax.set_yticks(list(ys))
    ax.set_yticklabels([r["country"] for r in data], fontweight="bold")
    ax.set_ylim(-0.6, n - 0.4)
    axt.set_ylim(-0.6, n - 0.4)
    ax.set_xlim(*(B_LIM if kind == "b" else C_LIM))
    ax.set_xlabel("Mean frequency of loneliness (0 = never, 10 = always)" if kind == "b"
                  else "Regression coefficient (versus reference category)")
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.spines["bottom"].set_linewidth(0.6)
    axt.axis("off")
    axt.set_title("Estimate (95% CI)", fontsize=9.5, pad=3)
    fig.savefig(path, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)


# --------------------------------------------------------------------------
# 3. Titles and legends
# --------------------------------------------------------------------------
def clean_var(v):
    return re.sub(r"\s*\(Ref:.*\)$", "", v).strip()


def ref_of(v):
    m = re.search(r"\(Ref:\s*(.+)\)$", v)
    return m.group(1).strip() if m else None


def sentence_case(s):
    return s[:1].lower() + s[1:] if not s[:2].isupper() else s


def title_for(num, kind, key):
    var, cat = key
    if kind == "b":
        return (f"Figure S{num}. Country-specific mean frequency of loneliness among "
                f"participants in the category {sentence_case(var)}: {cat}")
    ref = ref_of(var)
    if var.startswith("Race and ethnicity"):
        return (f"Figure S{num}. Country-specific associations of belonging to racial and ethnic "
                f"minority groups (versus the racial and ethnic plurality group) with frequency of loneliness")
    ref_txt = f" (reference: {ref})" if ref else ""
    return (f"Figure S{num}. Country-specific associations of {sentence_case(clean_var(var))}: "
            f"{cat}{ref_txt} with frequency of loneliness")


def missing_txt(data):
    miss = [c for c in COUNTRIES if c not in {r["country"] for r in data}]
    if not miss:
        return ""
    return (f" Countries not shown ({', '.join(miss)}) had no estimate for this category "
            f"(the variable was not assessed, the category had no respondents, or a different "
            f"reference category was used).")


LEG_B = ("Survey-weighted mean frequency of loneliness (0 = never, 10 = always) among participants "
         "in this category in each country. Black squares show country-specific means and horizontal "
         "lines show 95% confidence intervals (CIs); values are given to the right of the plot. "
         "Countries are ordered by mean. Estimates account for survey weights, clustering and "
         "stratification and were pooled across five multiply imputed datasets. The unit of analysis "
         f"is the individual survey respondent; N = {N_TOTAL} participants across 22 countries "
         "(country-specific sample sizes are given in Supplementary Tables S1a–S22a). Global P "
         "values for differences across categories (two-sided design-based Wald tests) are reported "
         "in Supplementary Tables S1b–S22b.")
LEG_C = ("Survey-weighted regression coefficients for this category versus the reference category, "
         "from country-specific linear regression models of frequency of loneliness (0 = never, "
         "10 = always) that included all childhood factors simultaneously. Black squares show "
         "coefficients and horizontal lines show 95% confidence intervals (CIs); values are given to "
         "the right of the plot. The vertical dashed line indicates no association; a 95% CI that "
         "excludes this line corresponds to a two-sided Wald test with P < 0.05, without adjustment "
         "for multiple comparisons. Countries are ordered by coefficient. Estimates account for survey "
         "weights, clustering and stratification and were pooled across five multiply imputed "
         "datasets. The unit of analysis is the individual survey respondent; "
         f"N = {N_TOTAL} participants across 22 countries (country-specific sample sizes are given in "
         "Supplementary Tables S1a–S22a). Exact global P values are reported in Supplementary "
         "Tables S1c–S22c.")

# --------------------------------------------------------------------------
# 4. Word document: title, image, legend; one figure per page
# --------------------------------------------------------------------------
out = docx.Document()
sec = out.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Mm(20))
st = out.styles["Normal"]
st.font.name = "Times New Roman"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.font.size = Pt(10)
st.paragraph_format.space_after = Pt(4)

for num, (kind, key, data) in enumerate(figs, start=1):
    png = os.path.join(OUT, "png", f"FigureS{num}.png")
    draw(kind, data, png)
    p = out.add_paragraph()
    if num > 1:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    p.add_run(title_for(num, kind, key)).bold = True
    pi = out.add_paragraph()
    pi.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pi.paragraph_format.keep_with_next = True
    pi.add_run().add_picture(png, width=Mm(FIG_W_MM))
    leg = (LEG_B if kind == "b" else LEG_C) + missing_txt(data)
    no_ci = [r["country"] for r in data if r["lo"] is None]
    if no_ci:
        leg += (f" For {', '.join(no_ci)}, very few participants were in this category, so no "
                f"standard error or 95% CI could be estimated (as in the corresponding Supplementary "
                f"Table); only the point estimate is shown.")
    if kind == "c" and key[0].lower().startswith("religious affiliation"):
        leg += (" Only countries in which no religion/atheist/agnostic was the reference category "
                "are shown.")
    out.add_paragraph(leg)

out.save(os.path.join(OUT, "SI_Figures_S1-S64.docx"))
print("done")
